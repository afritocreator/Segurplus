"""Test del Excel de salida contra valores calculados a mano, leyendo el
archivo generado con openpyxl (no solo que no explote)."""

from io import BytesIO

from openpyxl import load_workbook

from core.analisis.alertas import Alerta
from core.analisis.variacion import descomponer_variacion
from core.reportes.excel import generar_reporte_excel


def test_genera_excel_con_hojas_y_totales_correctos():
    d = descomponer_variacion(
        "abono_movil", cantidad_0=4, precio_0=2500, cantidad_1=6, precio_1=2800
    )
    alertas = [
        Alerta(tipo="salto_de_cantidad", severidad="media", mensaje="test", concepto="abono_movil")
    ]
    cuarentena = [("factura_rota.pdf", "la línea 1 no cierra")]

    buffer = BytesIO()
    resultado = generar_reporte_excel(
        servicio="telefonia",
        periodo_0="2026-07-01",
        periodo_1="2026-08-01",
        descomposiciones=[d],
        alertas=alertas,
        cuarentena=cuarentena,
        ruta_salida=buffer,
    )

    assert resultado is buffer
    wb = load_workbook(buffer)
    assert wb.sheetnames == ["Resumen", "Descomposición", "Alertas", "Cuarentena"]

    ws_resumen = wb["Resumen"]
    # total_0 = 4*2500 = 10000, total_1 = 6*2800 = 16800, variación = 6800
    assert ws_resumen.cell(row=3, column=2).value == 10000.0
    assert ws_resumen.cell(row=4, column=2).value == 16800.0
    assert ws_resumen.cell(row=5, column=2).value == 6800.0
    assert ws_resumen.cell(row=6, column=2).value == 1  # 1 concepto
    assert ws_resumen.cell(row=7, column=2).value == 1  # 1 alerta
    assert ws_resumen.cell(row=8, column=2).value == 1  # 1 en cuarentena

    ws_desc = wb["Descomposición"]
    assert ws_desc.cell(row=4, column=1).value == "abono_movil"
    assert ws_desc.cell(row=4, column=6).value == 5000.0  # efecto cantidad, calculado a mano
    assert ws_desc.cell(row=4, column=7).value == 1200.0  # efecto precio
    assert ws_desc.cell(row=4, column=9).value == 6800.0  # variación total

    ws_alertas = wb["Alertas"]
    assert ws_alertas.cell(row=4, column=2).value == "salto_de_cantidad"

    ws_cuarentena = wb["Cuarentena"]
    assert ws_cuarentena.cell(row=4, column=1).value == "factura_rota.pdf"


def test_sin_alertas_ni_cuarentena_dice_explicitamente_que_no_hay():
    buffer = BytesIO()
    generar_reporte_excel(
        servicio="energia",
        periodo_0="2026-07-01",
        periodo_1="2026-08-01",
        descomposiciones=[],
        alertas=[],
        cuarentena=[],
        ruta_salida=buffer,
    )
    wb = load_workbook(buffer)
    assert "Sin alertas" in wb["Alertas"].cell(row=4, column=1).value
    assert "No hay facturas en cuarentena" in wb["Cuarentena"].cell(row=4, column=1).value


def test_escribe_a_ruta_de_archivo(tmp_path):
    ruta = tmp_path / "sub" / "reporte.xlsx"
    resultado = generar_reporte_excel(
        servicio="gas",
        periodo_0="2026-07-01",
        periodo_1="2026-08-01",
        descomposiciones=[],
        alertas=[],
        cuarentena=[],
        ruta_salida=ruta,
    )
    assert resultado == ruta
    assert ruta.exists()
