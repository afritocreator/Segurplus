"""Exporta el resultado de una comparación de evolución (descomposición
precio/cantidad + alertas + cuarentena) a un Excel de varias hojas.

No calcula nada acá -- solo formatea. Todo el cálculo ya pasó por
core/analisis/ y sus tests con valor conocido. Mismo patrón de estilos que
core/reportes/excel.py de Consultora, para que el look sea consistente
entre las herramientas de la organización.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from core.analisis.alertas import Alerta
from core.analisis.variacion import DescomposicionVariacion

_TITULO_FONT = Font(bold=True, size=14, color="14324D")
_ENCABEZADO_FONT = Font(bold=True, color="FFFFFF")
_ENCABEZADO_FILL = PatternFill("solid", fgColor="14324D")
_ADVERTENCIA_FONT = Font(bold=True, italic=True, color="B00020")
_MONEDA = "#,##0.00"


def _encabezado(ws: Worksheet, fila: int, columnas: list[str]) -> None:
    for i, texto in enumerate(columnas, start=1):
        celda = ws.cell(row=fila, column=i, value=texto)
        celda.font = _ENCABEZADO_FONT
        celda.fill = _ENCABEZADO_FILL


def _hoja_descomposicion(wb: Workbook, descomposiciones: list[DescomposicionVariacion]) -> None:
    ws = wb.create_sheet("Descomposición")
    ws.cell(row=1, column=1, value="Descomposición precio × cantidad").font = _TITULO_FONT

    columnas = [
        "Concepto",
        "Cantidad (base)",
        "Precio (base)",
        "Cantidad (comparado)",
        "Precio (comparado)",
        "Efecto cantidad",
        "Efecto precio",
        "Efecto cruzado",
        "Variación total",
    ]
    _encabezado(ws, 3, columnas)

    fila = 4
    for d in descomposiciones:
        valores = [
            d.concepto,
            d.cantidad_0,
            d.precio_0,
            d.cantidad_1,
            d.precio_1,
            d.efecto_cantidad,
            d.efecto_precio,
            d.efecto_cruzado,
            d.variacion_total,
        ]
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            if col >= 2:
                celda.number_format = _MONEDA
        fila += 1

    ws.column_dimensions["A"].width = 28
    for col_letra in "BCDEFGHI":
        ws.column_dimensions[col_letra].width = 16


def _hoja_alertas(wb: Workbook, alertas: list[Alerta]) -> None:
    ws = wb.create_sheet("Alertas")
    ws.cell(row=1, column=1, value="Alertas").font = _TITULO_FONT
    _encabezado(ws, 3, ["Severidad", "Tipo", "Concepto", "Mensaje"])

    fila = 4
    for a in alertas:
        ws.cell(row=fila, column=1, value=a.severidad.upper())
        ws.cell(row=fila, column=2, value=a.tipo)
        ws.cell(row=fila, column=3, value=a.concepto or "")
        celda_mensaje = ws.cell(row=fila, column=4, value=a.mensaje)
        if a.severidad == "alta":
            celda_mensaje.font = _ADVERTENCIA_FONT
        fila += 1

    if not alertas:
        ws.cell(row=4, column=1, value="Sin alertas para esta comparación.")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 70


def _hoja_cuarentena(wb: Workbook, cuarentena: list[tuple[str, str]]) -> None:
    """`cuarentena`: lista de (ruta_pdf, motivos)."""
    ws = wb.create_sheet("Cuarentena")
    ws.cell(
        row=1, column=1, value="Facturas en cuarentena (no entraron al análisis)"
    ).font = _TITULO_FONT
    _encabezado(ws, 3, ["Archivo", "Motivo"])

    fila = 4
    for ruta, motivos in cuarentena:
        ws.cell(row=fila, column=1, value=ruta)
        celda_motivo = ws.cell(row=fila, column=2, value=motivos)
        celda_motivo.font = _ADVERTENCIA_FONT
        celda_motivo.alignment = Alignment(wrap_text=True)
        fila += 1

    if not cuarentena:
        ws.cell(row=4, column=1, value="No hay facturas en cuarentena.")

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 80


def generar_reporte_excel(
    *,
    servicio: str,
    periodo_0: str,
    periodo_1: str,
    descomposiciones: list[DescomposicionVariacion],
    alertas: list[Alerta],
    cuarentena: list[tuple[str, str]],
    ruta_salida: str | Path | BytesIO,
) -> Path | BytesIO:
    """Escribe el Excel de una comparación de evolución.

    `ruta_salida` acepta una ruta de archivo (str/Path) o un buffer en
    memoria (BytesIO — lo que usa el botón de descarga de Streamlit, sin
    tocar disco). Devuelve la ruta final, o el mismo buffer recibido, ya
    con el contenido escrito y el cursor al inicio."""
    wb = Workbook()

    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    ws_resumen.cell(
        row=1, column=1, value=f"Evolución de {servicio}: {periodo_0} → {periodo_1}"
    ).font = _TITULO_FONT

    total_0 = sum(d.total_0 for d in descomposiciones)
    total_1 = sum(d.total_1 for d in descomposiciones)
    filas_resumen: list[tuple[str, Any, str | None]] = [
        ("Total período base", total_0, _MONEDA),
        ("Total período comparado", total_1, _MONEDA),
        ("Variación total", total_1 - total_0, _MONEDA),
        ("Cantidad de conceptos", len(descomposiciones), None),
        ("Cantidad de alertas", len(alertas), None),
        ("Facturas en cuarentena", len(cuarentena), None),
    ]
    fila = 3
    for etiqueta, valor, formato in filas_resumen:
        ws_resumen.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
        celda_valor = ws_resumen.cell(row=fila, column=2, value=valor)
        if formato:
            celda_valor.number_format = formato
        fila += 1
    ws_resumen.column_dimensions["A"].width = 30
    ws_resumen.column_dimensions["B"].width = 20

    _hoja_descomposicion(wb, descomposiciones)
    _hoja_alertas(wb, alertas)
    _hoja_cuarentena(wb, cuarentena)

    if isinstance(ruta_salida, BytesIO):
        wb.save(ruta_salida)
        ruta_salida.seek(0)
        return ruta_salida

    ruta_salida = Path(ruta_salida)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta_salida)
    return ruta_salida
